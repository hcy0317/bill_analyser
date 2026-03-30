"""
WeChat Parser - 微信账单解析器

解析微信支付账单导出的 CSV 和 Excel (xlsx) 文件。
支持两种格式：
1. CSV格式（旧版导出）
2. XLSX格式（新版导出，从第17行开始为列名）

输出标准格式:
- date: 交易时间 (YYYY-MM-DD HH:MM:SS)
- amount: 金额 (支出为负, 收入为正)
- type: 类型 (收入/支出/转账/退款)
- description: 聚合描述
- source_account_id: 'wechat'
"""

import csv
from pathlib import Path
from typing import Any

import openpyxl

from ..utils.logger import log_method
from .base import ParserBase


class WeChatParser(ParserBase):
    """微信账单解析器"""

    # 解析器标识符
    PARSER_ID = "wechat"
    PARSER_NAME = "微信支付"

    def __init__(self):
        """初始化"""
        super().__init__()
        self.supported_extensions = [".csv", ".xlsx"]
        self.logger.info("微信账单解析器已初始化 [ID=%s], 支持格式: %s", self.PARSER_ID, self.supported_extensions)

    @log_method
    def can_parse(self, file_path: str) -> bool:
        """判断是否为微信账单"""
        file_ext = Path(file_path).suffix.lower()

        if file_ext == ".csv":
            return self._can_parse_csv(file_path)
        if file_ext == ".xlsx":
            return self._can_parse_xlsx(file_path)
        return False

    def _can_parse_csv(self, file_path: str) -> bool:
        """判断CSV文件是否为微信账单"""
        try:
            text, _encoding = self.read_text_with_fallback(file_path, ("utf-8", "utf-8-sig", "gbk", "gb18030"))
            first_lines = "\n".join(text.splitlines()[:8])

            if "微信支付账单" in first_lines:
                return True

            required_header_tokens = ["交易时间", "交易类型", "交易对方", "收/支", "金额(元)"]
            return all(token in first_lines for token in required_header_tokens)
        except Exception:  # pylint: disable=broad-except
            return False

    def _can_parse_xlsx(self, file_path: str) -> bool:
        """判断XLSX文件是否为微信账单"""
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            ws = wb.active
            if ws is None:
                wb.close()
                return False
            # 检查前几行是否包含微信账单特征
            for row in ws.iter_rows(max_row=5, values_only=True):
                if row[0] and "微信支付账单明细" in str(row[0]):
                    wb.close()
                    return True
            wb.close()
            return False
        except Exception:  # pylint: disable=broad-except
            return False

    @log_method
    def parse(self, file_path: str) -> list[dict[str, Any]]:
        """解析微信账单"""
        self.logger.info("[解析开始] 文件: %s", file_path)

        if not self.validate_file(file_path):
            return []

        file_ext = Path(file_path).suffix.lower()

        if file_ext == ".csv":
            return self._parse_csv(file_path)
        if file_ext == ".xlsx":
            return self._parse_xlsx(file_path)

        self.logger.error("不支持的文件格式: %s", file_ext)
        return []

    @log_method
    def _parse_csv(self, file_path: str) -> list[dict[str, Any]]:
        """解析CSV格式的微信账单"""
        bills = []

        try:
            with open(file_path, encoding="utf-8") as f:
                # 跳过头部说明行
                lines = f.readlines()
                data_start = 0

                for i, line in enumerate(lines):
                    if "交易时间" in line:
                        data_start = i
                        break

                if data_start == 0:
                    self.logger.warning("未找到数据起始行")
                    return []

                # 解析CSV数据
                reader = csv.DictReader(lines[data_start:])

                for row in reader:
                    try:
                        # 跳过空行和统计行
                        if not row.get("交易时间") or "总计" in str(row.get("交易时间")):
                            continue

                        bill = self._extract_bill_from_csv_row(row)
                        if bill:
                            bills.append(bill)

                    except Exception as e:  # pylint: disable=broad-except
                        self.logger.error("解析CSV行数据失败: %s", e)

            self.logger.info("[解析完成] CSV微信账单: %d 条", len(bills))

        except Exception as e:  # pylint: disable=broad-except
            self.logger.error("解析CSV微信账单失败: %s", e)
            return []

        return self.post_process(bills)

    @log_method
    def _parse_xlsx(self, file_path: str) -> list[dict[str, Any]]:
        """解析XLSX格式的微信账单"""
        bills = []

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            if ws is None:
                self.logger.error("未找到活动工作表")
                wb.close()
                return []

            # 找到列名行（通常在第17行左右）
            header_row = None
            header_row_idx = 0

            for idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                if row[0] == "交易时间":
                    header_row = row
                    header_row_idx = idx
                    self.logger.info("找到列名行，位于第 %d 行", idx)
                    break

            if not header_row:
                self.logger.error("未找到列名行")
                wb.close()
                return []

            # 创建列名到索引的映射
            column_map = {str(col): i for i, col in enumerate(header_row) if col}
            self.logger.debug("列名映射: %s", column_map)

            # 解析数据行
            for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
                try:
                    # 跳过空行
                    if not row[0]:
                        continue

                    # 跳过分隔线和统计行
                    if "----" in str(row[0]) or "注：" in str(row[0]) or "总计" in str(row[0]):
                        continue

                    bill = self._extract_bill_from_xlsx_row(row, column_map)
                    if bill:
                        bills.append(bill)

                except Exception as e:  # pylint: disable=broad-except
                    self.logger.error("解析XLSX行数据失败: %s, 行数据: %s", e, row)

            wb.close()
            self.logger.info("[解析完成] XLSX微信账单: %d 条", len(bills))

        except Exception as e:  # pylint: disable=broad-except
            self.logger.error("解析XLSX微信账单失败: %s", e)
            return []

        return self.post_process(bills)

    def _extract_bill_from_csv_row(self, row: dict[str, str]) -> dict[str, Any] | None:
        """从CSV行数据提取账单信息"""
        # 提取所有可能有用的字段
        return {
            "date": row.get("交易时间", ""),
            "type": row.get("收/支", ""),
            "counterparty": row.get("交易对方", ""),
            "description": row.get("商品", ""),
            "goods": row.get("商品", ""),
            "amount": row.get("金额(元)", "0"),
            "payment_method": row.get("支付方式", ""),
            "status": row.get("当前状态", ""),
            "transaction_id": row.get("交易单号", ""),
            "merchant_id": row.get("商户单号", ""),
            "remark": row.get("备注", ""),
            "original_category": row.get("交易类型", ""),
        }

    def _extract_bill_from_xlsx_row(self, row: tuple, column_map: dict[str, int]) -> dict[str, Any] | None:
        """从XLSX行数据提取账单信息"""
        try:
            # 获取列值的辅助函数
            def get_cell(col_name: str) -> str:
                idx = column_map.get(col_name)
                if idx is not None and idx < len(row):
                    value = row[idx]
                    return str(value) if value is not None else ""
                return ""

            # 提取金额并去除¥符号
            amount = get_cell("金额(元)")
            if amount.startswith("¥"):
                amount = amount[1:]

            return {
                "date": get_cell("交易时间"),
                "type": get_cell("收/支"),
                "counterparty": get_cell("交易对方"),
                "description": get_cell("商品"),
                "goods": get_cell("商品"),
                "amount": amount,
                "payment_method": get_cell("支付方式"),
                "status": get_cell("当前状态"),
                "transaction_id": get_cell("交易单号"),
                "merchant_id": get_cell("商户单号"),
                "remark": get_cell("备注"),
                "original_category": get_cell("交易类型"),
            }

        except Exception as e:  # pylint: disable=broad-except
            self.logger.error("提取XLSX账单信息失败: %s", e)
            return None
