from __future__ import annotations

from openpyxl import Workbook

from bill_analyser.parsers.wechat import WeChatParser
from tests.parser_test_support import build_wechat_xlsx, sample_path


def test_wechat_parser_handles_csv_sample() -> None:
    parser = WeChatParser()
    bills = parser.parse(str(sample_path("wechat_statement_sample.csv")))

    assert parser.can_parse(str(sample_path("wechat_statement_sample.csv"))) is True
    assert len(bills) == 7
    assert bills[0]["source_account_id"] == "wechat"
    assert any(bill["amount"] < 0 for bill in bills)
    assert any(bill["amount"] > 0 for bill in bills)


def test_wechat_parser_handles_generated_xlsx(tmp_path) -> None:
    parser = WeChatParser()
    xlsx_path = build_wechat_xlsx(tmp_path / "wechat_statement_sample.xlsx")

    bills = parser.parse(str(xlsx_path))

    assert parser.can_parse(str(xlsx_path)) is True
    assert len(bills) == 2
    assert bills[0]["description"] == "鲜肉包 | 测试早餐店 | 零钱 | 商户消费"
    assert bills[0]["amount"] == -5.0
    assert bills[1]["amount"] == 30.0


def test_wechat_parser_rejects_unknown_csv() -> None:
    parser = WeChatParser()

    assert parser.can_parse(str(sample_path("generic_statement_sample.csv"))) is False


def test_wechat_parser_handles_missing_headers_and_bad_workbooks(tmp_path) -> None:
    parser = WeChatParser()
    bad_csv = tmp_path / "bad_wechat.csv"
    bad_csv.write_text("备注,金额\n无效,1\n", encoding="utf-8")

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["不是微信账单"])
    bad_xlsx = tmp_path / "bad_wechat.xlsx"
    workbook.save(bad_xlsx)
    workbook.close()

    assert parser.parse(str(bad_csv)) == []
    assert parser.can_parse(str(bad_xlsx)) is False
    assert parser.parse(str(bad_xlsx)) == []


def test_wechat_parser_detects_signature_headers_and_unsupported_extension(tmp_path) -> None:
    parser = WeChatParser()
    signature_csv = tmp_path / "signature_wechat.csv"
    signature_csv.write_text(
        "交易时间,交易类型,交易对方,商品,收/支,金额(元)\n"
        "2026-01-15 08:01:00,商户消费,测试早餐店,鲜肉包,支出,¥5.00\n",
        encoding="utf-8-sig",
    )
    partial_csv = tmp_path / "partial_wechat.csv"
    partial_csv.write_text("交易时间,金额(元)\n2026-01-15 08:01:00,¥5.00\n", encoding="utf-8")

    unsupported_file = tmp_path / "wechat_statement.txt"
    unsupported_file.write_text("placeholder", encoding="utf-8")
    parser.supported_extensions.append(".txt")

    assert parser.can_parse(str(signature_csv)) is True
    assert parser.can_parse(str(partial_csv)) is False
    assert parser.parse(str(unsupported_file)) == []


def test_wechat_parser_handles_csv_row_errors_and_sheetless_workbooks(tmp_path, monkeypatch) -> None:
    parser = WeChatParser()
    valid_csv = tmp_path / "valid_wechat.csv"
    valid_csv.write_text(
        "微信支付账单\n"
        "交易时间,交易类型,交易对方,商品,收/支,金额(元),支付方式,当前状态,交易单号,商户单号,备注\n"
        "2026-01-15 08:01:00,商户消费,测试早餐店,鲜肉包,支出,¥5.00,零钱,支付成功,4200,WX-1,/\n",
        encoding="utf-8",
    )

    def explode_row(_row):
        raise RuntimeError("bad row")

    monkeypatch.setattr(parser, "_extract_bill_from_csv_row", explode_row)

    class SheetlessWorkbook:
        active = None

        def close(self) -> None:
            return None

    monkeypatch.setattr("bill_analyser.parsers.wechat.openpyxl.load_workbook", lambda *_args, **_kwargs: SheetlessWorkbook())

    assert parser.parse(str(valid_csv)) == []
    assert parser._can_parse_xlsx(str(tmp_path / "sheetless.xlsx")) is False  # pylint: disable=protected-access
    assert parser._parse_xlsx(str(tmp_path / "sheetless.xlsx")) == []  # pylint: disable=protected-access
    assert parser._extract_bill_from_xlsx_row(None, {"金额(元)": 0}) is None  # type: ignore[arg-type]  # pylint: disable=protected-access


def test_wechat_parser_covers_probe_failures_and_validate_short_circuit(tmp_path, monkeypatch) -> None:
    parser = WeChatParser()
    unsupported_file = tmp_path / "wechat_statement.txt"
    unsupported_file.write_text("placeholder", encoding="utf-8")

    assert parser.can_parse(str(unsupported_file)) is False

    monkeypatch.setattr(
        parser,
        "read_text_with_fallback",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("bad csv probe")),
    )
    assert parser._can_parse_csv("broken.csv") is False  # pylint: disable=protected-access

    monkeypatch.setattr(
        "bill_analyser.parsers.wechat.openpyxl.load_workbook",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("bad workbook")),
    )
    assert parser._can_parse_xlsx("broken.xlsx") is False  # pylint: disable=protected-access

    parser.validate_file = lambda _file_path: False  # type: ignore[method-assign]
    assert parser.parse(str(unsupported_file)) == []


def test_wechat_parser_skips_total_rows_and_handles_csv_io_failures(tmp_path, monkeypatch) -> None:
    parser = WeChatParser()
    csv_path = tmp_path / "wechat_totals.csv"
    csv_path.write_text(
        "微信支付账单\n"
        "交易时间,交易类型,交易对方,商品,收/支,金额(元),支付方式,当前状态,交易单号,商户单号,备注\n"
        "总计,统计项,全部,/,支出,¥0.00,零钱,支付成功,TOTAL,WX-TOTAL,/\n"
        "2026-01-15 08:01:00,商户消费,测试早餐店,鲜肉包,支出,¥5.00,零钱,支付成功,4200,WX-1,/\n",
        encoding="utf-8",
    )

    bills = parser.parse(str(csv_path))

    assert len(bills) == 1
    assert bills[0]["amount"] == -5.0

    monkeypatch.setattr("builtins.open", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("csv failed")))
    assert parser._parse_csv(str(csv_path)) == []  # pylint: disable=protected-access


def test_wechat_parser_xlsx_skips_empty_and_summary_rows(tmp_path) -> None:
    parser = WeChatParser()
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "账单"
    sheet.append(["微信支付账单明细"])
    sheet.append([""])
    sheet.append(
        ["交易时间", "交易类型", "交易对方", "商品", "收/支", "金额(元)", "支付方式", "当前状态", "交易单号", "商户单号", "备注"]
    )
    sheet.append([None])
    sheet.append(["---- 分隔线"])
    sheet.append(["注：这是说明"])
    sheet.append(["总计"])
    sheet.append(["2026-01-15 08:01:00", "商户消费", "测试早餐店", "鲜肉包", "支出", "¥5.00", "零钱", "支付成功", "4200", "WX-1", "/"])
    xlsx_path = tmp_path / "wechat_edge.xlsx"
    workbook.save(xlsx_path)
    workbook.close()

    bills = parser._parse_xlsx(str(xlsx_path))  # pylint: disable=protected-access
    extracted = parser._extract_bill_from_xlsx_row(("2026-01-15 08:01:00",), {"交易时间": 0})  # pylint: disable=protected-access

    assert len(bills) == 1
    assert bills[0]["amount"] == -5.0
    assert extracted is not None
    assert extracted["amount"] == ""
    assert extracted["payment_method"] == ""

    plain_amount_row = parser._extract_bill_from_xlsx_row(  # pylint: disable=protected-access
        ("2026-01-15 08:01:00", "支出", "测试早餐店", "鲜肉包", "支出", "5.00"),
        {"交易时间": 0, "收/支": 1, "交易对方": 2, "商品": 3, "金额(元)": 5},
    )
    assert plain_amount_row is not None
    assert plain_amount_row["amount"] == "5.00"


def test_wechat_parser_xlsx_handles_inner_and_outer_exceptions(tmp_path, monkeypatch) -> None:
    parser = WeChatParser()
    xlsx_path = build_wechat_xlsx(tmp_path / "wechat_statement_sample.xlsx")

    monkeypatch.setattr(
        parser,
        "_extract_bill_from_xlsx_row",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("bad xlsx row")),
    )
    assert parser._parse_xlsx(str(xlsx_path)) == []  # pylint: disable=protected-access

    monkeypatch.setattr(
        "bill_analyser.parsers.wechat.openpyxl.load_workbook",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("bad workbook")),
    )
    assert parser._parse_xlsx(str(xlsx_path)) == []  # pylint: disable=protected-access


def test_wechat_parser_skips_rows_when_extractors_return_none(tmp_path, monkeypatch) -> None:
    parser = WeChatParser()
    csv_path = tmp_path / "wechat_none_rows.csv"
    csv_path.write_text(
        "微信支付账单\n"
        "交易时间,交易类型,交易对方,商品,收/支,金额(元),支付方式,当前状态,交易单号,商户单号,备注\n"
        "2026-01-15 08:01:00,商户消费,测试早餐店,鲜肉包,支出,¥5.00,零钱,支付成功,4200,WX-1,/\n",
        encoding="utf-8",
    )

    monkeypatch.setattr(parser, "_extract_bill_from_csv_row", lambda _row: None)
    assert parser._parse_csv(str(csv_path)) == []  # pylint: disable=protected-access

    parser = WeChatParser()
    xlsx_path = build_wechat_xlsx(tmp_path / "wechat_none_rows.xlsx")
    monkeypatch.setattr(parser, "_extract_bill_from_xlsx_row", lambda *_args, **_kwargs: None)
    assert parser._parse_xlsx(str(xlsx_path)) == []  # pylint: disable=protected-access
