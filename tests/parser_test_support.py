from __future__ import annotations

from pathlib import Path
from textwrap import dedent

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SAMPLE_DIR = PROJECT_ROOT / 'tests' / 'fixtures' / 'import_samples'


def sample_path(name: str) -> Path:
    return SAMPLE_DIR / name


def _active_sheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    assert sheet is not None
    return sheet


def build_wechat_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.title = '账单'
    sheet.append(['微信支付账单明细'])
    sheet.append(['微信昵称：[pytest用户]'])
    sheet.append(['起始时间：[2026-01-01 00:00:00] 终止时间：[2026-01-31 23:59:59]'])
    sheet.append(['导出类型：[全部]'])
    sheet.append(['导出时间：[2026-03-29 12:00:00]'])
    for _ in range(11):
        sheet.append([''])
    sheet.append(
        ['交易时间', '交易类型', '交易对方', '商品', '收/支', '金额(元)', '支付方式', '当前状态', '交易单号', '商户单号', '备注']
    )
    sheet.append(
        ['2026-01-15 08:01:00', '商户消费', '测试早餐店', '鲜肉包', '支出', '¥5.00', '零钱', '支付成功', '4200000000000000000000000001', 'WX-PYTEST-0001', '/']
    )
    sheet.append(
        ['2026-01-16 09:31:00', '转账', '测试室友', '转账备注:微信转账', '收入', '¥30.00', '/', '已存入零钱', '100005000120260116000000000001', '/', '/']
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def build_icbc_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.title = '工行账单'
    sheet.append(['中国工商银行历史明细'])
    sheet.append(['账号', '交易日期', '收入/支出金额', '对方户名', '对方账号', '摘要', '交易流水号', '交易附言'])
    sheet.append(
        ['6222000000000000', '2026-01-10 08:00:00', -15.5, '测试早餐店', '6222000000000001', '早餐消费', 'ICBC-XLSX-0001', '早餐']
    )
    sheet.append(
        ['6222000000000000', '2026-01-11 19:30:00', 1200.0, '测试工资账户', '6222000000000002', '工资入账', 'ICBC-XLSX-0002', '工资']
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def build_icbc_html_xls(path: Path) -> Path:
        html = dedent(
                """
                <html>
                    <body>
                        <table>
                            <tr><td>中国工商银行</td></tr>
                            <tr>
                                <td>交易日期</td><td>收入/支出金额</td><td>对方户名</td><td>对方账号</td><td>摘要</td><td>交易流水号</td><td>交易附言</td>
                            </tr>
                            <tr>
                                <td>2026-01-10 08:00:00</td><td>-15.50</td><td>测试早餐店</td><td>6222000000000001</td><td>早餐消费</td><td>ICBC-HTML-0001</td><td>早餐</td>
                            </tr>
                            <tr>
                                <td>2026-01-11 19:30:00</td><td>1200.00</td><td>测试工资账户</td><td>6222000000000002</td><td>工资入账</td><td>ICBC-HTML-0002</td><td>工资</td>
                            </tr>
                        </table>
                    </body>
                </html>
                """
        ).strip()
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(html, encoding='utf-8')
        return path


def build_abc_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.title = '农行账单'
    sheet.append(['中国农业银行账户明细查询'])
    sheet.append([''])
    sheet.append(['交易日期', '交易时间', '收入金额', '支出金额', '对方户名', '交易用途', '交易摘要', '本次余额'])
    sheet.append(['2026-01-12', '08:15:00', '', '18.60', '测试早餐店', '早餐', '门店消费', '1000.00'])
    sheet.append(['2026-01-12', '18:20:00', '200.00', '', '测试报销账户', '差旅报销', '公司报销', '1200.00'])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def build_ccb_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.title = '建行账单'
    sheet.append(['中国建设银行'])
    sheet.append(['开户机构：测试支行'])
    sheet.append(['账　　号：6227000000000000'])
    sheet.append([''])
    sheet.append(['记账日', '交易日期', '交易时间', '摘要', '支出', '收入', '账户余额', '对方户名'])
    sheet.append(['20260112', '20260112', '081500', '早餐消费', 18.6, '', 1000.0, '测试早餐店'])
    sheet.append(['20260112', '20260112', '182000', '工资入账', '', 200.0, 1200.0, '测试工资账户'])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def build_cmbc_html_xls(path: Path) -> Path:
    html = dedent(
        """
        <html>
          <body>
            <table>
              <tr><td>中国民生银行股份有限公司</td></tr>
              <tr><td>个人账户对账单</td></tr>
              <tr>
                <td>交易时间</td><td>支出金额</td><td>存入金额</td><td>账户余额</td><td>对方账号</td><td>对方名称</td><td>对方开户行</td><td>摘要</td><td>交易方式</td>
              </tr>
              <tr>
                <td>20170101\t08:00:00</td><td>32.80</td><td></td><td>1000.00</td><td>6226000000000001</td><td>测试午餐店</td><td>民生银行</td><td>工作餐</td><td>手机银行</td>
              </tr>
              <tr>
                <td>20170102\t09:30:00</td><td></td><td>88.00</td><td>1088.00</td><td>6226000000000002</td><td>测试报销账户</td><td>民生银行</td><td>报销到账</td><td>网上银行</td>
              </tr>
            </table>
          </body>
        </html>
        """
    ).strip()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(html, encoding='utf-8')
    return path


def build_cmbc_xlsx(path: Path) -> Path:
    workbook = Workbook()
    sheet = _active_sheet(workbook)
    sheet.title = '民生账单'
    sheet.append(['中国民生银行股份有限公司'])
    sheet.append(['交易时间', '支出金额', '存入金额', '账户余额', '对方账号', '对方名称', '对方开户行', '摘要', '交易方式'])
    sheet.append(['20170101 08:00:00', '32.80', '', '1000.00', '6226000000000001', '测试午餐店', '民生银行', '工作餐', '手机银行'])
    sheet.append(['20170102 09:30:00', '', '88.00', '1088.00', '6226000000000002', '测试报销账户', '民生银行', '报销到账', '网上银行'])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path
