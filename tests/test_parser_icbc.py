"""Regression and edge-case tests for the ICBC parser."""

from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook

from bill_analyser.parsers.factory import ParserFactory
from bill_analyser.parsers.icbc import ICBCParser
from tests.parser_test_support import build_icbc_html_xls, build_icbc_xlsx, sample_path

# pylint: disable=missing-function-docstring,missing-class-docstring,too-few-public-methods


def test_icbc_parser_handles_csv_sample() -> None:
    parser = ICBCParser()
    bills = parser.parse(str(sample_path("icbc_statement_sample.csv")))

    assert parser.can_parse(str(sample_path("icbc_statement_sample.csv"))) is True
    assert len(bills) == 2
    assert bills[0]["source_account_id"] == "icbc"
    assert bills[0]["amount"] < 0
    assert bills[1]["amount"] > 0


def test_icbc_parser_handles_generated_xlsx(tmp_path) -> None:
    parser = ICBCParser()
    xlsx_path = build_icbc_xlsx(tmp_path / "icbc_statement_sample.xlsx")

    bills = parser.parse(str(xlsx_path))

    assert parser.can_parse(str(xlsx_path)) is True
    assert len(bills) == 2
    assert bills[0]["description"] == "早餐消费 | 测试早餐店 | 6222000000000001"
    assert bills[1]["amount"] == 1200.0


def test_icbc_internal_helpers_cover_content_detection(tmp_path) -> None:
    parser = ICBCParser()
    xlsx_path = build_icbc_xlsx(tmp_path / "icbc_statement_sample.xlsx")

    is_html, error = parser._detect_file_format(str(xlsx_path))  # pylint: disable=protected-access
    extracted = parser._extract_bill_from_excel_row(  # pylint: disable=protected-access
        ("6222000000000000", "2026-01-10 08:00:00", -15.5, "测试早餐店", "6222", "早餐消费", "流水1", "早餐"),
        {
            "交易日期": 1,
            "收入/支出金额": 2,
            "对方户名": 3,
            "对方账号": 4,
            "摘要": 5,
            "交易流水号": 6,
            "交易附言": 7,
        },
    )

    assert is_html is False
    assert error is None
    assert parser._is_icbc_content("收入/支出金额 对方账号名称 交易附言") is True  # pylint: disable=protected-access
    assert parser._is_icbc_content("民生银行 支出金额 存入金额") is False  # pylint: disable=protected-access
    assert extracted is not None
    assert extracted["type"] == "支出"


def test_icbc_parser_handles_html_and_content_based_detection() -> None:
    parser = ICBCParser()
    with tempfile.TemporaryDirectory(prefix="parser-neutral-") as temp_dir:
        neutral_root = Path(temp_dir)
        neutral_xlsx = build_icbc_xlsx(neutral_root / "statement.xlsx")
        html_xls = build_icbc_html_xls(neutral_root / "statement.xls")

        html_bills = parser.parse(str(html_xls))

        assert parser.can_parse(str(neutral_xlsx)) is True
        assert parser.can_parse(str(html_xls)) is True
        assert len(html_bills) == 2
        assert html_bills[0]["amount"] == -15.5


def test_icbc_parser_handles_bad_inputs_and_reader_fallbacks(tmp_path, monkeypatch) -> None:
    parser = ICBCParser()
    bad_csv = tmp_path / "bad_icbc.csv"
    bad_csv.write_text("无效内容\n", encoding="gbk")

    is_html, error = parser._detect_file_format(str(tmp_path / "missing.xls"))  # pylint: disable=protected-access
    monkeypatch.setattr(
        "bill_analyser.parsers.icbc.pd.read_excel",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("bad excel")),
    )

    assert parser._read_excel_content(str(tmp_path / "missing.xlsx")) == ""  # pylint: disable=protected-access
    assert parser.parse(str(bad_csv)) == []
    assert is_html is False
    assert error is not None


def test_icbc_parser_detects_neutral_utf8sig_csv_and_reads_html_content(tmp_path) -> None:
    parser = ICBCParser()
    neutral_csv = tmp_path / "statement.csv"
    neutral_csv.write_text(
        "中国工商银行历史明细,,,,,,\n"
        "交易日期,交易金额,对方户名,对方账号,摘要,交易流水号,收/支\n"
        "2026-01-10 08:00:00,-15.50,测试早餐店,6222000000000001,早餐消费,ICBC-0001,支出\n"
        "2026-01-11 19:30:00,1200.00,测试工资账户,6222000000000002,工资入账,ICBC-0002,收入\n",
        encoding="utf-8-sig",
    )
    html_xls = build_icbc_html_xls(tmp_path / "statement.xls")

    assert parser.can_parse(str(neutral_csv)) is True
    assert len(parser.parse(str(neutral_csv))) == 2
    assert "中国工商银行" in parser._read_html_content(str(html_xls))  # pylint: disable=protected-access


def test_icbc_parse_accepts_csv_when_header_is_first_line(monkeypatch) -> None:
    parser = ICBCParser()

    monkeypatch.setattr(parser, "validate_file", lambda _file_path: True)
    monkeypatch.setattr(
        parser,
        "read_lines_with_fallback",
        lambda _file_path: (
            [
                "交易日期,交易金额,收/支,对方户名,对方账号,摘要,交易流水号\n",
                "2026-01-03 08:00:00,-15.50,支出,测试早餐店,6222000000000001,早餐消费,ICBC-0003\n",
            ],
            "utf-8",
        ),
    )

    parsed_bills = parser.parse("header_first.csv")

    assert len(parsed_bills) == 1
    assert parsed_bills[0]["type"] == "支出"
    assert parsed_bills[0]["amount"] == -15.5


def test_icbc_factory_detects_neutral_csv_with_first_line_header(tmp_path) -> None:
    factory = ParserFactory()
    neutral_csv = tmp_path / "statement.csv"
    neutral_csv.write_text(
        "交易日期,交易金额,收/支,对方户名,对方账号,摘要,交易流水号\n"
        "2026-01-03 08:00:00,-15.50,支出,测试早餐店,6222000000000001,早餐消费,ICBC-0003\n",
        encoding="utf-8-sig",
    )

    parser = factory.get_parser(str(neutral_csv))
    parsed_bills = factory.parse(str(neutral_csv))

    assert parser is not None
    assert parser.PARSER_ID == "icbc"
    assert len(parsed_bills) == 1
    assert parsed_bills[0]["amount"] == -15.5


def test_icbc_factory_preserves_positive_income_without_explicit_type_columns(tmp_path) -> None:
    factory = ParserFactory()
    neutral_csv = tmp_path / "income_only.csv"
    neutral_csv.write_text(
        "交易日期,交易金额,对方户名,对方账号,摘要,交易流水号\n"
        "2026-01-03 08:00:00,1200.00,测试工资账户,6222000000000002,工资入账,ICBC-0004\n",
        encoding="utf-8-sig",
    )

    parsed_bills = factory.parse(str(neutral_csv))

    assert len(parsed_bills) == 1
    assert parsed_bills[0]["type"] == "收入"
    assert parsed_bills[0]["amount"] == 1200.0


def test_icbc_parse_html_xls_handles_empty_tables_missing_headers_and_bad_rows(monkeypatch) -> None:
    parser = ICBCParser()

    monkeypatch.setattr("bill_analyser.parsers.icbc.pd.read_html", lambda *_args, **_kwargs: [])
    assert not parser._parse_html_xls("empty.xls")  # pylint: disable=protected-access

    no_header_df = pd.DataFrame([["无效表头"], ["仍然无效"]])
    monkeypatch.setattr(
        "bill_analyser.parsers.icbc.pd.read_html",
        lambda *_args, **_kwargs: [no_header_df],
    )
    assert not parser._parse_html_xls("no_header.xls")  # pylint: disable=protected-access

    class BadStringValue:
        def __str__(self) -> str:
            raise RuntimeError("bad row payload")

    html_df = pd.DataFrame(
        [
            ["交易日期", "收入/支出金额", "对方户名", "对方账号", "摘要", "交易流水号"],
            [None, "10.00", "跳过空日期", "A001", "空日期", "L001"],
            ["2026-01-03 08:00:00", "bad", "坏金额", "A002", "坏金额", "L002"],
            ["2026-01-04 08:00:00", "0", "零金额", "A003", "零金额", "L003"],
            ["2026-01-05 08:00:00", "88.00", "坏摘要", "A004", BadStringValue(), "L004"],
            ["2026-01-06 08:00:00", "-12.50", "最终有效", "A005", "早餐消费", "L005"],
        ]
    )
    monkeypatch.setattr(
        "bill_analyser.parsers.icbc.pd.read_html",
        lambda *_args, **_kwargs: [html_df],
    )
    parsed_bills = parser._parse_html_xls("edge_cases.xls")  # pylint: disable=protected-access

    assert len(parsed_bills) == 1
    assert parsed_bills[0]["counterparty"] == "最终有效"
    assert parsed_bills[0]["amount"] == -12.5

    def _raise_html_error(*_args: Any, **_kwargs: Any) -> list[pd.DataFrame]:
        raise RuntimeError("broken html parser")

    monkeypatch.setattr("bill_analyser.parsers.icbc.pd.read_html", _raise_html_error)
    assert not parser._parse_html_xls("broken_html.xls")  # pylint: disable=protected-access


def test_icbc_parse_excel_covers_missing_headers_invalid_rows_and_top_level_errors(
    tmp_path,
    monkeypatch,
) -> None:
    parser = ICBCParser()

    missing_header_path = tmp_path / "missing_header.xlsx"
    missing_header_workbook = Workbook()
    missing_header_sheet = missing_header_workbook.active
    assert missing_header_sheet is not None
    missing_header_sheet.append(["首行说明"])
    missing_header_workbook.save(missing_header_path)
    missing_header_workbook.close()

    assert not parser._parse_excel(str(missing_header_path))  # pylint: disable=protected-access

    edge_path = tmp_path / "edge_cases.xlsx"
    edge_workbook = Workbook()
    edge_sheet = edge_workbook.active
    assert edge_sheet is not None
    edge_sheet.append(["说明行"])
    edge_sheet.append(
        ["账号", "交易日期", "收入/支出金额", "对方户名", "对方账号", "摘要", "交易流水号", "交易附言"]
    )
    edge_sheet.append(["6222", None, -10.0, "空日期", "A001", "空日期", "L001", "附言"])
    edge_sheet.append(["6222", "2026-01-01 08:00:00", "bad", "坏金额", "A002", "坏金额", "L002", "附言"])
    edge_sheet.append(["6222", "2026-01-02 08:00:00", "0", "零金额", "A003", "零金额", "L003", "附言"])
    edge_sheet.append(["6222", "2026-01-03 08:00:00", -15.5, "最终有效", "A004", "早餐消费", "L004", "附言"])
    edge_workbook.save(edge_path)
    edge_workbook.close()

    parsed_bills = parser._parse_excel(str(edge_path))  # pylint: disable=protected-access
    assert len(parsed_bills) == 1
    assert parsed_bills[0]["counterparty"] == "最终有效"
    assert parsed_bills[0]["amount"] == -15.5

    def _raise_workbook_error(*_args: Any, **_kwargs: Any) -> Workbook:
        raise RuntimeError("broken workbook")

    monkeypatch.setattr("bill_analyser.parsers.icbc.openpyxl.load_workbook", _raise_workbook_error)
    assert not parser._parse_excel("broken.xlsx")  # pylint: disable=protected-access


def test_icbc_can_parse_and_read_helpers_cover_failure_branches(monkeypatch) -> None:
    parser = ICBCParser()

    assert parser.can_parse("statement.txt") is False

    monkeypatch.setattr(
        parser,
        "read_text_with_fallback",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("bad csv probe")),
    )
    assert parser._can_parse_csv("broken.csv") is False  # pylint: disable=protected-access

    monkeypatch.setattr(
        parser,
        "read_text_with_fallback",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(UnicodeDecodeError("utf-8", b"x", 0, 1, "bad html")),
    )
    assert parser._read_html_content("broken.xls") == ""  # pylint: disable=protected-access

    monkeypatch.setattr(parser, "_detect_file_format", lambda _file_path: (False, "cannot probe"))
    assert parser._can_parse_excel("broken.xlsx") is False  # pylint: disable=protected-access

    monkeypatch.setattr(parser, "_detect_file_format", lambda _file_path: (False, None))
    monkeypatch.setattr(parser, "_read_excel_content", lambda _file_path: "")
    assert parser._can_parse_excel("empty.xlsx") is False  # pylint: disable=protected-access

    monkeypatch.setattr(
        parser,
        "_detect_file_format",
        lambda _file_path: (_ for _ in ()).throw(RuntimeError("explode")),
    )
    assert parser._can_parse_excel("explode.xlsx") is False  # pylint: disable=protected-access


def test_icbc_parse_covers_validate_short_circuit_and_excel_fallback(monkeypatch) -> None:
    parser = ICBCParser()
    monkeypatch.setattr(parser, "validate_file", lambda _file_path: False)
    assert parser.parse("ignored.csv") == []

    parser = ICBCParser()
    fallback_bills = [{"marker": "fallback"}]
    monkeypatch.setattr(parser, "validate_file", lambda _file_path: True)
    monkeypatch.setattr("builtins.open", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("bad header read")))
    monkeypatch.setattr(parser, "_parse_excel", lambda _file_path: fallback_bills)
    assert parser.parse("broken.xlsx") == fallback_bills

    parser = ICBCParser()
    monkeypatch.setattr(parser, "validate_file", lambda _file_path: True)
    assert parser.parse("unsupported.txt") == []


def test_icbc_parse_csv_and_excel_cover_outer_and_inner_exceptions(tmp_path, monkeypatch) -> None:
    parser = ICBCParser()
    monkeypatch.setattr(
        parser,
        "read_lines_with_fallback",
        lambda _file_path: (_ for _ in ()).throw(RuntimeError("bad csv read")),
    )
    assert parser._parse_csv("broken.csv") == []  # pylint: disable=protected-access

    class SheetlessWorkbook:
        active = None

        def close(self) -> None:
            return None

    monkeypatch.setattr("bill_analyser.parsers.icbc.openpyxl.load_workbook", lambda *_args, **_kwargs: SheetlessWorkbook())
    assert parser._parse_excel("sheetless.xlsx") == []  # pylint: disable=protected-access

    parser = ICBCParser()
    xlsx_path = build_icbc_xlsx(tmp_path / "icbc_statement_sample.xlsx")
    monkeypatch.setattr(
        ICBCParser,
        "_extract_bill_from_excel_row",
        lambda self, *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("bad excel row")),
    )
    assert parser._parse_excel(str(xlsx_path)) == []  # pylint: disable=protected-access


def test_icbc_parse_html_and_csv_cover_remaining_skip_and_exception_branches(tmp_path, monkeypatch) -> None:
    parser = ICBCParser()

    monkeypatch.setattr(
        "bill_analyser.parsers.icbc.pd.read_html",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(UnicodeDecodeError("utf-8", b"x", 0, 1, "all bad")),
    )
    assert parser._parse_html_xls("all_bad_encodings.xls") == []  # pylint: disable=protected-access

    monkeypatch.setattr(
        parser,
        "read_lines_with_fallback",
        lambda _file_path: (
            [
                "中国工商银行历史明细\n",
                "交易日期,交易金额,对方户名,对方账号,摘要,交易流水号\n",
                ",-15.50,空日期,6222000000000001,早餐消费,ICBC-EMPTY\n",
                "2026-01-10 08:00:00,-15.50,测试早餐店,6222000000000001,早餐消费,ICBC-0001\n",
            ],
            "utf-8",
        ),
    )
    csv_bills = parser._parse_csv("skip_none_rows.csv")  # pylint: disable=protected-access

    assert len(csv_bills) == 1
    assert csv_bills[0]["counterparty"] == "测试早餐店"


def test_icbc_parse_excel_covers_empty_header_columns_and_row_exception_logging(tmp_path, monkeypatch) -> None:
    parser = ICBCParser()
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["说明行"])
    sheet.append(["账号", "交易日期", "收入/支出金额", None, "对方户名", "对方账号", "摘要", "交易流水号", "交易附言"])
    sheet.append(["bad", "2026-01-02 08:00:00", "-8.80", None, "坏行商户", "A000", "坏行", "L000", "附言"])
    sheet.append(["6222", "2026-01-03 08:00:00", "-15.50", None, "测试早餐店", "A001", "早餐消费", "L001", "附言"])
    xlsx_path = tmp_path / "icbc_empty_header.xlsx"
    workbook.save(xlsx_path)
    workbook.close()

    class BadExcelRowValue:
        def __str__(self) -> str:
            raise RuntimeError("bad excel row value")

    original_extract = ICBCParser._extract_bill_from_excel_row

    def _extract_then_raise(self, row, column_map):
        if row and row[0] == "bad":
            raise RuntimeError("bad row in parser loop")
        return original_extract(self, row, column_map)

    monkeypatch.setattr(ICBCParser, "_extract_bill_from_excel_row", _extract_then_raise)
    bills = parser._parse_excel(str(xlsx_path))  # pylint: disable=protected-access

    assert len(bills) == 1
    assert bills[0]["counterparty"] == "测试早餐店"

    assert original_extract(  # pylint: disable=protected-access
        parser,
        ("2026-01-01 08:00:00", BadExcelRowValue()),
        {"交易日期": 0, "收入/支出金额": 1},
    ) is None


def test_icbc_helper_extractors_cover_zero_missing_retry_and_exception_paths(monkeypatch) -> None:
    parser = ICBCParser()

    assert parser._extract_bill_from_html_row({"交易日期": "2026-01-01", "收入/支出金额": "0.0"}) is None  # pylint: disable=protected-access
    assert parser._extract_bill_from_csv_row({"交易金额": "10"}) is None  # pylint: disable=protected-access

    income_row = parser._extract_bill_from_csv_row(  # pylint: disable=protected-access
        {"交易日期": "2026-01-01 08:00:00", "交易金额": "-10", "借贷标志": "贷"}
    )
    assert income_row is not None
    assert income_row["type"] == "收入"

    class BadStringValue:
        def __str__(self) -> str:
            raise RuntimeError("bad string conversion")

    assert parser._extract_bill_from_csv_row({"交易日期": "2026-01-01 08:00:00", "交易金额": BadStringValue()}) is None  # pylint: disable=protected-access
    assert parser._extract_bill_from_excel_row(("only",), {}) is None  # pylint: disable=protected-access
    assert parser._extract_bill_from_excel_row(  # pylint: disable=protected-access
        ("2026-01-01 08:00:00", "0.0"),
        {"交易日期": 0, "收入/支出金额": 1},
    ) is None
    assert parser._extract_bill_from_html_row({"交易日期": BadStringValue(), "收入/支出金额": "18.5"}) is None  # pylint: disable=protected-access
    assert parser._extract_bill_from_excel_row((BadStringValue(),), {"交易日期": 0}) is None  # pylint: disable=protected-access

    html_df = pd.DataFrame(
        [
            ["交易日期", "收入/支出金额", "对方户名", "对方账号", "摘要", "交易流水号"],
            ["2026-01-03 08:00:00", "-15.50", "测试早餐店", "6222", "早餐消费", "ICBC-0005"],
        ]
    )
    read_html_calls: list[str | None] = []

    def _read_html_with_retry(*_args, **kwargs) -> list[pd.DataFrame]:
        encoding = kwargs.get("encoding")
        read_html_calls.append(encoding)
        if len(read_html_calls) == 1:
            raise UnicodeDecodeError(str(encoding or "utf-8"), b"x", 0, 1, "bad html encoding")
        return [html_df]

    monkeypatch.setattr("bill_analyser.parsers.icbc.pd.read_html", _read_html_with_retry)
    html_bills = parser._parse_html_xls("retry.xls")  # pylint: disable=protected-access

    assert len(html_bills) == 1
    assert html_bills[0]["amount"] == -15.5
    assert len(read_html_calls) >= 2
