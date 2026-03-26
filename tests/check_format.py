"""临时脚本：检查账单文件格式"""
import openpyxl

# 检查工商银行xlsx文件
file_path = 'src/bills/工商银行历史明细（申请单号：25082520010041279887）.xlsx'

try:
    wb = openpyxl.load_workbook(file_path, read_only=True)
    ws = wb.active
    
    print("=== 工商银行账单结构 ===")
    print(f"总行数: {ws.max_row}")
    print(f"总列数: {ws.max_column}")
    print("\n前10行内容:")
    for i, row in enumerate(ws.iter_rows(max_row=10, values_only=True), 1):
        print(f"{i}: {row}")
        
    wb.close()
except Exception as e:
    print(f"错误: {e}")
